import json
from collections import OrderedDict
from copy import deepcopy
from pathlib import Path

try:
    import openpyxl
except ImportError as exc:
    raise SystemExit(
        "openpyxl is required to read dynamic-checking Excel reports. "
        "Install it or run this script in the project environment."
    ) from exc


ERROR_CATEGORIES = OrderedDict({
    "API Hallucination": {
        "no-member",
        "import-error",
        "no-name-in-module",
    },
    "API Parameters Misuse": {
        "unexpected-keyword-arg",
        "no-value-for-parameter",
        "too-many-function-args",
        "redundant-keyword-arg",
        "not-callable",
        "bad-super-call",
    },
})


class NormalizeDeeperStatistics:

    def __init__(
        self,
        static_report_root,
        dynamic_report_root,
        output_root,
        model_ids,
        step_id,
        experiment_id,
        output_suffix,
        digits=None,
    ):
        self.static_report_root = Path(static_report_root)
        self.dynamic_report_root = Path(dynamic_report_root)
        self.output_root = Path(output_root)
        self.model_ids = model_ids
        self.step_id = step_id
        self.experiment_id = experiment_id
        self.output_suffix = output_suffix
        self.digits = digits
        self.summary_sheet = "summary1" if step_id == "step_three" else "summary"
        self.valid_file_columns = ("valid_files", "valid_pyfiles")
        self.reserved_keys = {"source_file", "summary", "normalization"}

    @staticmethod
    def parse_count(value):
        if isinstance(value, int):
            return value
        if isinstance(value, float):
            return int(value)
        if isinstance(value, str):
            return int(value.split("(")[0])
        raise ValueError(f"Unsupported count value: {value}")

    def count_errors(self, error_counts):
        statistics = {
            "categories": OrderedDict(
                (category, {"total": 0, "issues": OrderedDict()})
                for category in ERROR_CATEGORIES
            ),
            "uncategorized": OrderedDict(),
        }

        for issue, raw_count in error_counts.items():
            count = self.parse_count(raw_count)
            matched_category = None

            for category, issues in ERROR_CATEGORIES.items():
                if issue in issues:
                    matched_category = category
                    break

            if matched_category is None:
                statistics["uncategorized"][issue] = count
                continue

            statistics["categories"][matched_category]["issues"][issue] = count
            statistics["categories"][matched_category]["total"] += count

        statistics["categories"] = OrderedDict(
            sorted(
                statistics["categories"].items(),
                key=lambda item: item[1]["total"],
                reverse=True
            )
        )
        statistics["uncategorized"] = OrderedDict(
            sorted(statistics["uncategorized"].items(), key=lambda item: item[1], reverse=True)
        )
        statistics["total_categorized_errors"] = sum(
            category_counts["total"]
            for category_counts in statistics["categories"].values()
        )
        statistics["total_uncategorized_errors"] = sum(statistics["uncategorized"].values())
        statistics["total_errors"] = (
            statistics["total_categorized_errors"] + statistics["total_uncategorized_errors"]
        )

        return statistics

    def load_model_statistics(self, model_id):
        report_dir = self.static_report_root / model_id / self.step_id
        deeper_path = report_dir / f"{model_id}_deeper_statistics.json"
        if deeper_path.exists():
            with deeper_path.open("r", encoding="utf-8") as file:
                return json.load(file), deeper_path, "deeper_statistics"

        messages_path = report_dir / f"{model_id}_messages_statistics.json"
        if not messages_path.exists():
            raise FileNotFoundError(f"Missing static deeper/messages statistics: {report_dir}")

        with messages_path.open("r", encoding="utf-8") as file:
            messages_data = json.load(file)

        model_statistics = {
            "source_file": str(messages_path),
            "summary": self.count_errors(messages_data.get("summary", {}).get("error", {})),
        }
        for dll, dll_data in messages_data.items():
            if dll in self.reserved_keys or not isinstance(dll_data, dict):
                continue

            experiment_data = dll_data.get(self.experiment_id, {})
            if isinstance(experiment_data, dict):
                model_statistics[dll] = {
                    self.experiment_id: self.count_errors(experiment_data.get("error", {}))
                }

        return model_statistics, messages_path, "messages_statistics"

    def read_valid_files(self, workbook_path):
        workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
        if self.summary_sheet not in workbook.sheetnames:
            raise ValueError(f"Workbook has no {self.summary_sheet!r} sheet: {workbook_path}")

        worksheet = workbook[self.summary_sheet]
        rows = list(worksheet.iter_rows(values_only=True))
        header_index, headers = None, {}

        for index, row in enumerate(rows):
            headers = {
                str(cell).strip().lower(): position
                for position, cell in enumerate(row)
                if cell is not None
            }
            if "benchmark" in headers and any(column in headers for column in self.valid_file_columns):
                header_index = index
                break

        if header_index is None:
            raise ValueError(f"Cannot find valid-files summary header: {workbook_path}")

        benchmark_column = headers["benchmark"]
        valid_column_name = next(column for column in self.valid_file_columns if column in headers)
        valid_column = headers[valid_column_name]
        data_rows = rows[header_index + 1:]
        matching_rows = [
            row
            for row in data_rows
            if len(row) > benchmark_column and str(row[benchmark_column]).strip() == self.step_id
        ]

        if not matching_rows:
            non_empty_rows = [row for row in data_rows if any(cell is not None for cell in row)]
            if len(non_empty_rows) == 1:
                matching_rows = non_empty_rows
            else:
                raise ValueError(f"Cannot find benchmark {self.step_id!r}: {workbook_path}")

        value = matching_rows[0][valid_column]
        denominator = float(value)
        if denominator <= 0:
            raise ValueError(f"Denominator must be positive in {workbook_path}, got {value!r}")

        return denominator

    def normalize_statistics(self, statistics, denominator):
        normalized = deepcopy(statistics)

        for category_data in normalized.get("categories", {}).values():
            category_data["total"] = self.divide(category_data.get("total", 0), denominator)
            category_data["issues"] = {
                issue: self.divide(count, denominator)
                for issue, count in category_data.get("issues", {}).items()
            }

        normalized["uncategorized"] = {
            issue: self.divide(count, denominator)
            for issue, count in normalized.get("uncategorized", {}).items()
        }

        for total_key in (
            "total_categorized_errors",
            "total_uncategorized_errors",
            "total_errors",
        ):
            if total_key in normalized:
                normalized[total_key] = self.divide(normalized[total_key], denominator)

        return normalized

    def divide(self, value, denominator):
        result = float(value) / denominator
        if self.digits is not None:
            return round(result, self.digits)
        return round(result, 4)

    def normalize_model(self, model_id):
        model_data, source_path, source_type = self.load_model_statistics(model_id)
        leaf_denominators = {}

        for dll, dll_data in model_data.items():
            if dll in self.reserved_keys or not isinstance(dll_data, dict):
                continue

            for experiment_id, experiment_data in dll_data.items():
                if not isinstance(experiment_data, dict):
                    continue

                workbook_path = (
                    self.dynamic_report_root
                    / model_id
                    / dll
                    / "dynamic_checking.xlsx"
                )
                if not workbook_path.exists():
                    raise FileNotFoundError(f"Missing dynamic report workbook: {workbook_path}")

                leaf_denominators[(dll, experiment_id)] = self.read_valid_files(workbook_path)

        summary_denominator = sum(leaf_denominators.values())
        normalized = {
            "source_file": str(source_path),
            "normalization": {
                "formula": "original_count / valid_files",
                "step_id": self.step_id,
                "source_type": source_type,
                "summary_sheet": self.summary_sheet,
                "valid_file_columns": list(self.valid_file_columns),
                "dynamic_report_root": str(self.dynamic_report_root),
                "summary_scope": "static-report-leaves-only",
                "summary_valid_files": summary_denominator,
                "leaf_valid_files": {
                    f"{dll}/{experiment_id}": denominator
                    for (dll, experiment_id), denominator in sorted(leaf_denominators.items())
                },
            },
        }

        if "summary" in model_data:
            normalized["summary"] = self.normalize_statistics(
                model_data["summary"],
                summary_denominator,
            )

        for dll, dll_data in model_data.items():
            if dll in self.reserved_keys or not isinstance(dll_data, dict):
                continue

            normalized[dll] = {}
            for experiment_id, experiment_data in dll_data.items():
                if not isinstance(experiment_data, dict):
                    continue

                denominator = leaf_denominators[(dll, experiment_id)]
                normalized[dll][experiment_id] = self.normalize_statistics(
                    experiment_data,
                    denominator,
                )

        return normalized

    def main(self):
        for model_id in self.model_ids:
            report = self.normalize_model(model_id)
            output_path = (
                self.output_root
                / model_id
                / self.step_id
                / f"{model_id}{self.output_suffix}"
            )
            output_path.parent.mkdir(parents=True, exist_ok=True)
            with output_path.open("w", encoding="utf-8") as file:
                json.dump(report, file, indent=2, ensure_ascii=False)
            print(f"Wrote {output_path}")


if __name__ == "__main__":
    CURRENT_DIR = Path(__file__).resolve().parent
    EVALUATION_DIR = CURRENT_DIR.parent

    STATIC_REPORT_ROOT = CURRENT_DIR / "report"
    DYNAMIC_REPORT_ROOT = EVALUATION_DIR / "dynamic_checking" / "report"
    OUTPUT_ROOT = CURRENT_DIR / "report"
    MODEL_ID_LIST = [
        "gemma_3_4b_it",
        "gemma_3_12b_it",
        "gemma_3_27b_it",
        "ministral_3_3b_instruct_2512",
        "ministral_3_14b_instruct_2512",
        "mistral_small_3_2_24b_instruct_2506"
    ]
    STEP_ID = "step_three"
    EXPERIMENT_ID = "models_0701"
    OUTPUT_SUFFIX = "_deeper_statistics_per_valid_file.json"
    DIGITS = None

    NormalizeDeeperStatistics(
        static_report_root=STATIC_REPORT_ROOT,
        dynamic_report_root=DYNAMIC_REPORT_ROOT,
        output_root=OUTPUT_ROOT,
        model_ids=MODEL_ID_LIST,
        step_id=STEP_ID,
        experiment_id=EXPERIMENT_ID,
        output_suffix=OUTPUT_SUFFIX,
        digits=DIGITS,
    ).main()
